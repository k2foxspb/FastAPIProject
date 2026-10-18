import asyncio
import calendar
import io
from datetime import date, datetime, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.dependencies import get_async_db
from app.core.auth import get_current_production_user
from app.core.production import (
    can_view_timesheet,
    can_edit_timesheet,
    can_review_production,
    can_edit_production_entry,
    find_nearest_supervisors,
    find_all_supervisors,
    get_subordinates,
)
from app.core.roles import get_table_role_label
from app.api.routers.notifications import manager as notification_manager
from app.core.fcm import send_fcm_notification
from app.models.users import User as UserModel
from app.models.structure import Workshop as WorkshopModel
from app.models.timesheet import (
    TimesheetEntry as TimesheetEntryModel,
    TimesheetEntryApproval as TimesheetEntryApprovalModel,
    TimesheetAuditLog as TimesheetAuditLogModel,
    Product as ProductModel,
    ProductionEntry as ProductionEntryModel,
    ProductionEntryItem as ProductionEntryItemModel,
    ProductionAuditLog as ProductionAuditLogModel,
    PRODUCTION_STATUS_DRAFT,
    PRODUCTION_STATUS_PENDING,
    PRODUCTION_STATUS_APPROVED,
    PRODUCTION_STATUS_REJECTED,
    TIMESHEET_APPROVAL_STATUS_NOTIFIED,
)
from app.schemas.timesheet import (
    TimesheetEntryUpsert,
    TimesheetEntry as TimesheetEntrySchema,
    TimesheetEntryDetail as TimesheetEntryDetailSchema,
    TimesheetEntryReject,
    TimesheetApprovalRound,
    Product as ProductSchema,
    ProductCreate,
    ProductUpdate,
    ProductionEntryUpsert,
    ProductionEntryReject,
    ProductionEntry as ProductionEntrySchema,
    ProductionEntryDetail as ProductionEntryDetailSchema,
    DayTimesheetInfo,
    SubordinateSummary,
    SubordinateDay,
    EmployeeBrief,
)

router = APIRouter(prefix="/timesheet", tags=["timesheet"])


def _user_name(user: UserModel | None) -> str:
    if user is None:
        return ""
    return f"{user.first_name or ''} {user.last_name or ''}".strip() or user.email or user.phone_number or f"#{user.id}"


async def _get_user_or_404(db: AsyncSession, user_id: int) -> UserModel:
    result = await db.execute(select(UserModel).where(UserModel.id == user_id))
    user = result.scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    return user


async def _load_timesheet_entry(db: AsyncSession, entry_id: int) -> TimesheetEntryModel:
    result = await db.execute(
        select(TimesheetEntryModel)
        .options(
            selectinload(TimesheetEntryModel.user),
            selectinload(TimesheetEntryModel.reviewer),
            selectinload(TimesheetEntryModel.approvals),
        )
        .where(TimesheetEntryModel.id == entry_id)
    )
    entry = result.scalar_one_or_none()
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Timesheet entry not found")
    return entry


async def _load_timesheet_entry_detail(db: AsyncSession, entry_id: int) -> TimesheetEntryModel:
    result = await db.execute(
        select(TimesheetEntryModel)
        .options(
            selectinload(TimesheetEntryModel.user),
            selectinload(TimesheetEntryModel.reviewer),
            selectinload(TimesheetEntryModel.approvals).selectinload(TimesheetEntryApprovalModel.approver),
            selectinload(TimesheetEntryModel.logs).selectinload(TimesheetAuditLogModel.actor),
        )
        .where(TimesheetEntryModel.id == entry_id)
    )
    entry = result.scalar_one_or_none()
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Timesheet entry not found")
    return entry


async def _resolve_required_approvers(
    db: AsyncSession, target: UserModel, extra_approver_user_ids: list[int] | None = None,
) -> list[UserModel]:
    """
    Определяет обязательный набор согласующих для изменения часов сотрудника
    target - ближайший вышестоящий руководитель (всегда) плюс дополнительные
    руководители, выбранные сотрудником галочкой "на согласование"
    (extra_approver_user_ids), отфильтрованные через всю цепочку
    руководителей, чтобы нельзя было указать произвольного пользователя. Все
    они должны согласовать изменение - если хотя бы один отклонит, изменение
    считается отклонённым целиком.
    """
    nearest = await find_nearest_supervisors(db, target)
    approvers = {s.id: s for s in nearest}
    if extra_approver_user_ids:
        all_supervisors = await find_all_supervisors(db, target)
        selected_ids = set(extra_approver_user_ids)
        for s in all_supervisors:
            if s.id in selected_ids:
                approvers[s.id] = s
    return list(approvers.values())


async def _resolve_notify_only_recipients(
    db: AsyncSession, target: UserModel, extra_notify_user_ids: list[int] | None, exclude_ids: set[int],
) -> list[UserModel]:
    """
    Определяет руководителей, которых нужно просто уведомить об изменении
    часов галочкой "уведомление" (без требования согласования) -
    отфильтрованные через всю цепочку руководителей. Руководители, уже
    попавшие в обязательные согласующие (exclude_ids), исключаются, чтобы не
    создавать для одного и того же человека дублирующие записи.
    """
    if not extra_notify_user_ids:
        return []
    all_supervisors = await find_all_supervisors(db, target)
    selected_ids = set(extra_notify_user_ids)
    return [s for s in all_supervisors if s.id in selected_ids and s.id not in exclude_ids]


def _timesheet_entry_to_schema(entry: TimesheetEntryModel) -> TimesheetEntrySchema:
    schema = TimesheetEntrySchema.model_validate(entry)
    pending = [
        a.approver for a in entry.approvals
        if a.round == entry.approval_round and a.status == PRODUCTION_STATUS_PENDING and a.approver is not None
    ]
    schema.pending_approvers = pending
    return schema


def _build_timesheet_history(entry: TimesheetEntryModel) -> list[TimesheetApprovalRound]:
    """
    Строит историю карточки дня по раундам - каждое событие отправки на
    согласование (или правки часов руководителем) формирует отдельный блок,
    внутри которого перечислено, кто и как согласовал/отклонил.
    """
    logs = sorted(entry.logs, key=lambda l: (l.round, l.created_at))
    rounds = []
    for log in logs:
        approvals = sorted(
            (a for a in entry.approvals if a.round == log.round),
            key=lambda a: a.id,
        )
        rounds.append(TimesheetApprovalRound.model_validate({
            "round": log.round,
            "action": log.action,
            "actor_id": log.actor_id,
            "actor": log.actor,
            "created_at": log.created_at,
            "approvals": approvals,
        }))
    return rounds


def _timesheet_entry_to_detail_schema(entry: TimesheetEntryModel) -> TimesheetEntryDetailSchema:
    schema = TimesheetEntryDetailSchema.model_validate(entry)
    pending = [
        a.approver for a in entry.approvals
        if a.round == entry.approval_round and a.status == PRODUCTION_STATUS_PENDING and a.approver is not None
    ]
    schema.pending_approvers = pending
    schema.history = _build_timesheet_history(entry)
    return schema


def _notify_supervisors_about_timesheet(
    entry: TimesheetEntryModel,
    supervisors: list[UserModel],
    resubmit: bool = False,
    notify_only: list[UserModel] | None = None,
):
    """
    Уведомляет всех обязательных согласующих (см. _resolve_required_approvers)
    об изменении часов, ожидающем согласования, а также руководителей,
    отмеченных только для уведомления (см. _resolve_notify_only_recipients) -
    им не нужно ничего согласовывать, они просто информируются об изменении.
    """
    employee_name = _user_name(entry.user)
    date_str = entry.date.isoformat()
    body = f"{employee_name} - {date_str}: {entry.hours} ч (+{entry.overtime_hours} сверхурочных)"

    def _send(recipients: list[UserModel], message_type: str, title: str, requires_action: bool):
        if not recipients:
            return
        message = {
            "type": message_type,
            "entry_id": entry.id,
            "employee_id": entry.user_id,
            "employee_name": employee_name,
            "date": date_str,
            "hours": str(entry.hours),
            "overtime_hours": str(entry.overtime_hours),
            "requires_action": requires_action,
        }
        for recipient in recipients:
            asyncio.create_task(notification_manager.send_personal_message(message, recipient.id))
            if recipient.fcm_token:
                asyncio.create_task(send_fcm_notification(
                    token=recipient.fcm_token,
                    title=title,
                    body=body,
                    data=message,
                    sender_id=entry.user_id,
                ))

    _send(
        supervisors,
        "timesheet_entry_resubmitted" if resubmit else "timesheet_entry_submitted",
        "Повторная отправка табеля на согласование" if resubmit else "Изменение табеля на согласовании",
        requires_action=True,
    )
    _send(
        notify_only or [],
        "timesheet_entry_notified",
        "Изменение табеля (уведомление)",
        requires_action=False,
    )


def _notify_employee_about_timesheet_review(entry: TimesheetEntryModel, approved: bool):
    """Уведомляет сотрудника о результате согласования изменения часов."""
    message = {
        "type": "timesheet_entry_approved" if approved else "timesheet_entry_rejected",
        "entry_id": entry.id,
        "date": entry.date.isoformat(),
        "comment": entry.reviewer_comment,
    }
    asyncio.create_task(notification_manager.send_personal_message(message, entry.user_id))
    if entry.user and entry.user.fcm_token:
        title = "Изменение табеля подтверждено" if approved else "Изменение табеля отклонено"
        body = f"За {entry.date.isoformat()}" if approved else (entry.reviewer_comment or f"За {entry.date.isoformat()}")
        asyncio.create_task(send_fcm_notification(
            token=entry.user.fcm_token,
            title=title,
            body=body,
            data=message,
            sender_id=entry.reviewer_id,
        ))


def _month_bounds(year: int, month: int) -> tuple[date, date]:
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def _default_base_hours(user: UserModel | None, day: date) -> Decimal:
    """
    Часы по умолчанию для дня, за который ещё нет записи в табеле - берутся
    из графика работы пользователя (work_schedule_hours), заданного им при
    регистрации/в профиле. По выходным (суббота, воскресенье) по умолчанию 0.
    """
    if day.weekday() >= 5:
        return Decimal("0")
    if user is None or user.work_schedule_hours is None:
        return Decimal("0")
    return Decimal(str(user.work_schedule_hours))


_MONTH_NAMES_RU = [
    "", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]


def _compute_month_totals(year: int, month: int, row_data: dict) -> tuple[Decimal, Decimal, Decimal, Decimal]:
    """
    Считает по одному сотруднику за месяц: сумму основного времени, сумму
    сверхурочного времени, сумму часов, отработанных в выходные дни, и общий
    итог (основное + сверхурочное). Текстовые отметки (например больничный)
    в подсчёт часов не включаются.
    """
    last_day = calendar.monthrange(year, month)[1]
    hours_by_date = row_data["hours_by_date"]
    base_sum = Decimal("0")
    overtime_sum = Decimal("0")
    weekend_sum = Decimal("0")
    for d in range(1, last_day + 1):
        cur_date = date(year, month, d)
        default_base = _default_base_hours(row_data.get("user"), cur_date)
        base_hours, overtime_hours, text_value = hours_by_date.get(
            cur_date, (default_base, Decimal("0"), None)
        )
        if text_value:
            continue
        base_sum += base_hours
        overtime_sum += overtime_hours
        if cur_date.weekday() in (5, 6):
            weekend_sum += base_hours + overtime_hours
    return base_sum, overtime_sum, weekend_sum, base_sum + overtime_sum


async def _get_workshop_label(db: AsyncSession, workshop_id: int | None) -> str | None:
    """Возвращает подпись "Цех № {id} «{название}»" для шапки Excel-выгрузки табеля."""
    if workshop_id is None:
        return None
    result = await db.execute(select(WorkshopModel).where(WorkshopModel.id == workshop_id))
    workshop = result.scalar_one_or_none()
    if not workshop:
        return None
    return f"Цех № {workshop.id} «{workshop.name}»"


def _write_export_summary_header(ws, workshop_label: str | None, year: int, month: int, rows: list[dict]) -> int:
    """
    Пишет в начале листа шапку выгрузки табеля: номер и название цеха, месяц
    и год, а также суммарные по всем сотрудникам из `rows` часы - общее
    отработанное время, отдельно основное, отдельно сверхурочное и отдельно
    за выходные. Возвращает номер строки, с которой должна начинаться сама
    таблица табеля.
    """
    from openpyxl.styles import Font

    total_base = Decimal("0")
    total_overtime = Decimal("0")
    total_weekend = Decimal("0")
    for row_data in rows:
        base_sum, overtime_sum, weekend_sum, _ = _compute_month_totals(year, month, row_data)
        total_base += base_sum
        total_overtime += overtime_sum
        total_weekend += weekend_sum
    total_all = total_base + total_overtime

    row = 1
    if workshop_label:
        ws.cell(row=row, column=1, value=workshop_label).font = Font(bold=True, size=13)
        row += 1
    ws.cell(
        row=row, column=1,
        value=f"Табель учёта рабочего времени за {_MONTH_NAMES_RU[month]} {year} года",
    ).font = Font(bold=True, size=12)
    row += 1
    ws.cell(row=row, column=1, value=f"Общее отработанное время всех работников: {float(total_all)} ч.").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value=f"в т.ч. основное время: {float(total_base)} ч.")
    row += 1
    ws.cell(row=row, column=1, value=f"в т.ч. сверхурочное время: {float(total_overtime)} ч.")
    row += 1
    ws.cell(row=row, column=1, value=f"в т.ч. время, отработанное в выходные: {float(total_weekend)} ч.")
    row += 2  # пустая строка-разделитель перед таблицей
    return row


def _fill_timesheet_grid_sheet(ws, year: int, month: int, rows: list[dict], start_row: int = 1) -> None:
    """
    Заполняет лист в формате классического табеля учёта рабочего времени:
    по дням месяца с разбивкой каждой строки сотрудника на "основное" и
    "сверхурочное" время (значения берутся напрямую из табеля, т.к. они
    вводятся сотрудником/руководителем явно, а не высчитываются по норме),
    подытогами за 1-ю половину месяца и за весь месяц, количеством часов,
    отработанных в выходные дни, общим количеством отработанных часов за
    месяц (основное + сверхурочное) и общим количеством сданных приборов
    в конце строки.

    rows: список словарей с ключами:
      number, tab_number, grade, fio,
      hours_by_date (dict[date, tuple[Decimal, Decimal, str | None]] -
        (основное, сверхурочное, текстовая отметка вместо часов)),
      products_total (Decimal)

    start_row: номер строки, с которой начинаются заголовки таблицы (по
      умолчанию 1) - используется, когда перед таблицей на листе уже
      выведена шапка с итогами (см. _write_export_summary_header).
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    weekend_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    last_day = calendar.monthrange(year, month)[1]
    days = list(range(1, last_day + 1))
    half = 15 if last_day >= 15 else last_day

    headers = ["№", "Таб. номер", "Разряд", "Фамилия, имя, отчество", "Время"]
    day_col_index: dict[int, int] = {}
    for d in days:
        headers.append(str(d))
        day_col_index[d] = len(headers)
        if d == half:
            headers.append("За 1-ю половину\n(дни)")
            half_days_col = len(headers)
            headers.append("За 1-ю половину\n(часы)")
            half_hours_col = len(headers)
    headers.append("За месяц\n(дни)")
    month_days_col = len(headers)
    headers.append("За месяц\n(часы)")
    month_hours_col = len(headers)
    headers.append("Выходные\n(часы)")
    weekend_col = len(headers)
    headers.append("Итого отработано\n(часов)")
    total_hours_col = len(headers)
    headers.append("Приборы, всего")
    products_col = len(headers)

    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=idx, value=title)
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = thin_border
        cell.fill = header_fill

    weekend_days = {d for d in days if date(year, month, d).weekday() in (5, 6)}
    for d in weekend_days:
        ws.cell(row=start_row, column=day_col_index[d]).fill = weekend_fill

    row_num = start_row + 1
    for row_data in rows:
        hours_by_date = row_data["hours_by_date"]

        # Общее количество отработанных часов за месяц (основное + сверхурочное
        # суммарно) считается один раз на сотрудника и выводится только в
        # первой (основной) подстроке - аналогично "Приборы, всего".
        _, _, _, total_month_hours = _compute_month_totals(year, month, row_data)

        for label in ("основное", "сверхурочное"):
            is_first = label == "основное"
            ws.cell(row=row_num, column=1, value=row_data["number"] if is_first else None).border = thin_border
            ws.cell(row=row_num, column=2, value=row_data["tab_number"] if is_first else None).border = thin_border
            ws.cell(row=row_num, column=3, value=row_data["grade"] if is_first else None).border = thin_border
            ws.cell(row=row_num, column=4, value=row_data["fio"] if is_first else None).border = thin_border
            ws.cell(row=row_num, column=5, value=label).border = thin_border

            half_days_count = 0
            half_hours_sum = Decimal("0")
            month_days_count = 0
            month_hours_sum = Decimal("0")
            weekend_hours_sum = Decimal("0")

            for d in days:
                cur_date = date(year, month, d)
                default_base = _default_base_hours(row_data.get("user"), cur_date)
                base_hours, overtime_hours, text_value = hours_by_date.get(
                    cur_date, (default_base, Decimal("0"), None)
                )
                value = base_hours if is_first else overtime_hours

                cell = ws.cell(row=row_num, column=day_col_index[d])
                cell.border = thin_border
                cell.alignment = center
                if d in weekend_days:
                    cell.fill = weekend_fill
                if text_value and is_first:
                    # Текстовая отметка (например больничный) показывается вместо числа часов
                    # и не учитывается в итогах.
                    cell.value = text_value
                elif value > 0:
                    cell.value = float(value)
                    month_days_count += 1
                    month_hours_sum += value
                    if d <= half:
                        half_days_count += 1
                        half_hours_sum += value
                    if d in weekend_days:
                        weekend_hours_sum += value

            ws.cell(row=row_num, column=half_days_col, value=half_days_count).border = thin_border
            ws.cell(row=row_num, column=half_hours_col, value=float(half_hours_sum)).border = thin_border
            ws.cell(row=row_num, column=month_days_col, value=month_days_count).border = thin_border
            ws.cell(row=row_num, column=month_hours_col, value=float(month_hours_sum)).border = thin_border
            # Часы, отработанные в выходные дни (а не просто их количество).
            ws.cell(row=row_num, column=weekend_col, value=float(weekend_hours_sum)).border = thin_border
            ws.cell(
                row=row_num, column=total_hours_col,
                value=float(total_month_hours) if is_first else None,
            ).border = thin_border
            ws.cell(
                row=row_num, column=products_col,
                value=float(row_data["products_total"]) if is_first else None,
            ).border = thin_border

            row_num += 1

    ws.column_dimensions["D"].width = 28
    for d in days:
        ws.column_dimensions[ws.cell(row=start_row, column=day_col_index[d]).column_letter].width = 6
    ws.freeze_panes = f"F{start_row + 1}"


# ------------------------------------------------------------------
# Табель рабочего времени
# ------------------------------------------------------------------

@router.get("/{user_id}", response_model=list[TimesheetEntrySchema])
async def get_user_timesheet(
    user_id: int,
    year: int,
    month: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """Возвращает табель пользователя за месяц (правило 1 - доступ по иерархии)."""
    target = await _get_user_or_404(db, user_id)
    if not can_view_timesheet(current_user, target):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to this timesheet")

    start, end = _month_bounds(year, month)
    result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id == user_id,
            TimesheetEntryModel.date >= start,
            TimesheetEntryModel.date <= end,
        )
    )
    return result.scalars().all()


@router.get("/{user_id}/supervisors", response_model=list[EmployeeBrief])
async def get_timesheet_supervisors(
    user_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Возвращает всю цепочку вышестоящих руководителей пользователя - список
    для чекбоксов на экране табеля, чтобы сотрудник мог дополнительно
    выбрать, кого из вышестоящих руководителей (помимо ближайшего, который
    уведомляется всегда) уведомить об изменении часов.
    """
    target = await _get_user_or_404(db, user_id)
    if not can_edit_timesheet(current_user, target):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to this timesheet")
    return await find_all_supervisors(db, target)


@router.get("/entries/pending", response_model=list[TimesheetEntrySchema])
async def list_pending_timesheet_entries(
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Изменения часов, ожидающие СОБСТВЕННОГО согласования текущего
    пользователя - то есть он назначен одним из обязательных согласующих в
    текущем раунде (см. _resolve_required_approvers) и ещё не принял
    решение. Если выбрано несколько руководителей, каждый видит запись в
    своём списке, пока не согласует хотя бы свою часть.
    """
    result = await db.execute(
        select(TimesheetEntryModel)
        .options(selectinload(TimesheetEntryModel.user), selectinload(TimesheetEntryModel.approvals))
        .where(TimesheetEntryModel.status == PRODUCTION_STATUS_PENDING)
        .order_by(TimesheetEntryModel.submitted_at.asc())
    )
    entries = result.scalars().all()
    visible = [
        e for e in entries
        if can_review_production(current_user, e.user)
        and any(
            a.round == e.approval_round and a.approver_id == current_user.id and a.status == PRODUCTION_STATUS_PENDING
            for a in e.approvals
        )
    ]
    return [_timesheet_entry_to_schema(e) for e in visible]


@router.put("/{user_id}", response_model=TimesheetEntrySchema)
async def upsert_timesheet_day(
    user_id: int,
    data: TimesheetEntryUpsert,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Создаёт/обновляет количество часов за день (правило 1 - доступ по
    иерархии). Если часы правит сам сотрудник - изменение уходит на
    согласование всем обязательным согласующим сразу (статус "pending") -
    ближайшему руководителю и дополнительным, выбранным галочками; все они
    должны согласовать, отказ любого отклоняет изменение целиком. Если часы
    правит вышестоящий пользователь - изменение считается сразу
    согласованным (статус "approved").
    """
    target = await _get_user_or_404(db, user_id)
    if not can_edit_timesheet(current_user, target):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to edit this timesheet")

    is_self_edit = current_user.id == target.id

    result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id == user_id,
            TimesheetEntryModel.date == data.date,
        )
    )
    entry = result.scalar_one_or_none()
    was_rejected = entry is not None and entry.status == PRODUCTION_STATUS_REJECTED

    if entry is None:
        entry = TimesheetEntryModel(
            user_id=user_id,
            date=data.date,
            created_by_id=current_user.id,
            approval_round=0,
        )
        db.add(entry)

    entry.hours = data.hours
    entry.overtime_hours = data.overtime_hours
    entry.text_value = data.text_value
    entry.comment = data.comment
    entry.updated_by_id = current_user.id

    required_approvers: list[UserModel] = []
    notify_only_recipients: list[UserModel] = []
    entry.approval_round = (entry.approval_round or 0) + 1

    if is_self_edit:
        required_approvers = await _resolve_required_approvers(db, target, data.extra_approver_user_ids)
        notify_only_recipients = await _resolve_notify_only_recipients(
            db, target, data.extra_notify_user_ids, exclude_ids={a.id for a in required_approvers},
        )
        entry.status = PRODUCTION_STATUS_PENDING
        entry.submitted_at = datetime.now(timezone.utc)
        entry.reviewer_comment = None
        entry.reviewed_at = None
        entry.reviewer_id = None
    else:
        entry.status = PRODUCTION_STATUS_APPROVED
        entry.reviewer_id = current_user.id
        entry.reviewed_at = datetime.now(timezone.utc)
        entry.reviewer_comment = None

    await db.flush()

    if is_self_edit:
        db.add(TimesheetAuditLogModel(
            entry_id=entry.id,
            round=entry.approval_round,
            actor_id=current_user.id,
            action="resubmitted" if was_rejected else "submitted",
        ))
        for approver in required_approvers:
            db.add(TimesheetEntryApprovalModel(
                entry_id=entry.id,
                round=entry.approval_round,
                approver_id=approver.id,
                status=PRODUCTION_STATUS_PENDING,
            ))
        for recipient in notify_only_recipients:
            db.add(TimesheetEntryApprovalModel(
                entry_id=entry.id,
                round=entry.approval_round,
                approver_id=recipient.id,
                status=TIMESHEET_APPROVAL_STATUS_NOTIFIED,
                decided_at=datetime.now(timezone.utc),
            ))
    else:
        db.add(TimesheetAuditLogModel(
            entry_id=entry.id,
            round=entry.approval_round,
            actor_id=current_user.id,
            action="updated_by_supervisor",
        ))

    await db.commit()

    entry = await _load_timesheet_entry(db, entry.id)

    if is_self_edit:
        _notify_supervisors_about_timesheet(
            entry, required_approvers, resubmit=was_rejected, notify_only=notify_only_recipients,
        )

    return _timesheet_entry_to_schema(entry)


@router.get("/entry/{entry_id}", response_model=TimesheetEntryDetailSchema)
async def get_timesheet_entry_detail(
    entry_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Карточка дня табеля с полной историей согласования часов внизу,
    сгруппированной по каждой отправке на согласование (кто отправил и кто
    как согласовал/отклонил в рамках этой отправки).
    """
    entry = await _load_timesheet_entry_detail(db, entry_id)
    if not can_view_timesheet(current_user, entry.user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    return _timesheet_entry_to_detail_schema(entry)


@router.post("/entry/{entry_id}/approve", response_model=TimesheetEntrySchema)
async def approve_timesheet_entry(
    entry_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Подтверждает свою часть согласования изменения часов. Если выбрано
    несколько обязательных согласующих, запись переходит в "approved"
    только когда согласуют ВСЕ из них (правило 11 - только вышестоящий
    пользователь, не сам сотрудник, и только тот, кто назначен согласующим
    в текущем раунде).
    """
    entry = await _load_timesheet_entry(db, entry_id)
    if not can_review_production(current_user, entry.user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if entry.status != PRODUCTION_STATUS_PENDING:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Entry is not pending approval")

    approval = next(
        (a for a in entry.approvals if a.round == entry.approval_round and a.approver_id == current_user.id
         and a.status == PRODUCTION_STATUS_PENDING),
        None,
    )
    if approval is None:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You are not assigned to approve this entry")

    approval.status = PRODUCTION_STATUS_APPROVED
    approval.decided_at = datetime.now(timezone.utc)
    approval.comment = None
    await db.flush()

    fully_approved = not any(
        a.round == entry.approval_round and a.status == PRODUCTION_STATUS_PENDING for a in entry.approvals
    )
    if fully_approved:
        entry.status = PRODUCTION_STATUS_APPROVED
        entry.reviewer_id = current_user.id
        entry.reviewed_at = datetime.now(timezone.utc)
        entry.reviewer_comment = None
    await db.commit()

    entry = await _load_timesheet_entry(db, entry.id)
    if fully_approved:
        _notify_employee_about_timesheet_review(entry, approved=True)
    return _timesheet_entry_to_schema(entry)


@router.post("/entry/{entry_id}/reject", response_model=TimesheetEntrySchema)
async def reject_timesheet_entry(
    entry_id: int,
    data: TimesheetEntryReject,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Отклоняет изменение часов - комментарий обязателен (правило 11). Отказ
    хотя бы одного из обязательных согласующих отклоняет запись целиком,
    независимо от решения остальных.
    """
    entry = await _load_timesheet_entry(db, entry_id)
    if not can_review_production(current_user, entry.user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if entry.status != PRODUCTION_STATUS_PENDING:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Entry is not pending approval")

    approval = next(
        (a for a in entry.approvals if a.round == entry.approval_round and a.approver_id == current_user.id
         and a.status == PRODUCTION_STATUS_PENDING),
        None,
    )
    if approval is None:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You are not assigned to approve this entry")

    approval.status = PRODUCTION_STATUS_REJECTED
    approval.decided_at = datetime.now(timezone.utc)
    approval.comment = data.comment

    entry.status = PRODUCTION_STATUS_REJECTED
    entry.reviewer_id = current_user.id
    entry.reviewed_at = datetime.now(timezone.utc)
    entry.reviewer_comment = data.comment
    await db.commit()

    entry = await _load_timesheet_entry(db, entry.id)
    _notify_employee_about_timesheet_review(entry, approved=False)
    return _timesheet_entry_to_schema(entry)


# ------------------------------------------------------------------
# Справочник продукции (CRUD)
# ------------------------------------------------------------------

@router.get("/products/list", response_model=list[ProductSchema])
async def list_products(
    include_inactive: bool = False,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    query = select(ProductModel)
    if not include_inactive:
        query = query.where(ProductModel.is_active == True)
    query = query.order_by(ProductModel.name.asc())
    result = await db.execute(query)
    return result.scalars().all()


@router.post("/products", response_model=ProductSchema, status_code=status.HTTP_201_CREATED)
async def create_product(
    data: ProductCreate,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    product = ProductModel(name=data.name, code=data.code, unit=data.unit, is_active=data.is_active)
    db.add(product)
    await db.commit()
    await db.refresh(product)
    return product


@router.put("/products/{product_id}", response_model=ProductSchema)
async def update_product(
    product_id: int,
    data: ProductUpdate,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    result = await db.execute(select(ProductModel).where(ProductModel.id == product_id))
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Product not found")

    for field in ("name", "code", "unit", "is_active"):
        value = getattr(data, field)
        if value is not None:
            setattr(product, field, value)

    await db.commit()
    await db.refresh(product)
    return product


@router.delete("/products/{product_id}")
async def delete_product(
    product_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """Деактивирует продукцию (мягкое удаление, чтобы не потерять историю сданных позиций)."""
    result = await db.execute(select(ProductModel).where(ProductModel.id == product_id))
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Product not found")
    product.is_active = False
    await db.commit()
    return {"status": "ok"}


# ------------------------------------------------------------------
# Записи о сданной продукции
# ------------------------------------------------------------------

async def _load_entry(db: AsyncSession, entry_id: int) -> ProductionEntryModel:
    result = await db.execute(
        select(ProductionEntryModel)
        .options(
            selectinload(ProductionEntryModel.items).selectinload(ProductionEntryItemModel.product),
            selectinload(ProductionEntryModel.user),
            selectinload(ProductionEntryModel.reviewer),
            selectinload(ProductionEntryModel.logs).selectinload(ProductionAuditLogModel.actor),
        )
        .where(ProductionEntryModel.id == entry_id)
    )
    entry = result.scalar_one_or_none()
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Production entry not found")
    return entry


async def _get_hours_for_date(db: AsyncSession, user: UserModel, day: date) -> tuple[Decimal, Decimal, str | None]:
    """
    Возвращает (основное, сверхурочное, текстовая отметка) из табеля за день.
    Если записи в табеле ещё нет, основное время по умолчанию берётся из
    графика работы пользователя (work_schedule_hours).
    """
    result = await db.execute(
        select(
            TimesheetEntryModel.hours, TimesheetEntryModel.overtime_hours, TimesheetEntryModel.text_value
        ).where(
            TimesheetEntryModel.user_id == user.id, TimesheetEntryModel.date == day
        )
    )
    row = result.one_or_none()
    if row is None:
        return _default_base_hours(user, day), Decimal("0"), None
    base_hours, overtime_hours, text_value = row
    return base_hours or Decimal("0"), overtime_hours or Decimal("0"), text_value


def _entry_to_schema(
    entry: ProductionEntryModel,
    current_user: UserModel,
    base_hours: Decimal | None = None,
    overtime_hours: Decimal | None = None,
    text_value: str | None = None,
) -> ProductionEntrySchema:
    schema = ProductionEntrySchema.model_validate(entry)
    schema.base_hours = base_hours
    schema.overtime_hours = overtime_hours
    schema.hours = (base_hours + overtime_hours) if base_hours is not None and overtime_hours is not None else None
    schema.text_value = text_value
    schema.can_edit = can_edit_production_entry(current_user, entry.user, entry.status)
    return schema


async def _notify_supervisors(db: AsyncSession, entry: ProductionEntryModel, hours: Decimal, resubmit: bool = False):
    supervisors = await find_nearest_supervisors(db, entry.user)
    if not supervisors:
        return

    items_summary = ", ".join(
        f"{item.product.name if item.product else item.product_id} x{item.quantity}" for item in entry.items
    )
    employee_name = _user_name(entry.user)
    message = {
        "type": "production_entry_resubmitted" if resubmit else "production_entry_submitted",
        "entry_id": entry.id,
        "employee_id": entry.user_id,
        "employee_name": employee_name,
        "date": entry.date.isoformat(),
        "items": items_summary,
        "hours": str(hours),
    }
    title = "Повторная отправка на согласование" if resubmit else "Новая сдача продукции на согласование"
    body = f"{employee_name} - {entry.date.isoformat()}: {items_summary or 'без позиций'} ({hours} ч)"

    for supervisor in supervisors:
        asyncio.create_task(notification_manager.send_personal_message(message, supervisor.id))
        if supervisor.fcm_token:
            asyncio.create_task(send_fcm_notification(
                token=supervisor.fcm_token,
                title=title,
                body=body,
                data=message,
                sender_id=entry.user_id,
            ))


@router.get("/subordinates/summary", response_model=list[SubordinateSummary])
async def get_subordinates_summary(
    year: int,
    month: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Сводная таблица по всем подчинённым текущего пользователя за месяц: часы
    из табеля по каждому дню и статус сданной продукции - чтобы видеть всех
    подчинённых сразу и переходить к редактированию часов/согласованию.
    Начальник цеха и выше видит в этой таблице и свою собственную строку.
    """
    subordinates = await get_subordinates(db, current_user, include_self=True)
    if not subordinates:
        return []

    sub_ids = [u.id for u in subordinates]
    start, end = _month_bounds(year, month)

    hours_result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id.in_(sub_ids),
            TimesheetEntryModel.date >= start,
            TimesheetEntryModel.date <= end,
        )
    )
    hours_by_key: dict[tuple[int, date], TimesheetEntryModel] = {
        (e.user_id, e.date): e for e in hours_result.scalars().all()
    }

    entries_result = await db.execute(
        select(ProductionEntryModel).where(
            ProductionEntryModel.user_id.in_(sub_ids),
            ProductionEntryModel.date >= start,
            ProductionEntryModel.date <= end,
        )
    )
    entries_by_key: dict[tuple[int, date], ProductionEntryModel] = {
        (e.user_id, e.date): e for e in entries_result.scalars().all()
    }

    summaries = []
    for user in subordinates:
        days = []
        total_hours = Decimal("0")
        pending_count = 0
        pending_timesheet_count = 0
        for day_num in range(1, end.day + 1):
            day = date(year, month, day_num)
            timesheet_entry = hours_by_key.get((user.id, day))
            if timesheet_entry is not None:
                base_hours = timesheet_entry.hours
                overtime_hours = timesheet_entry.overtime_hours
                text_value = timesheet_entry.text_value
            else:
                base_hours, overtime_hours, text_value = _default_base_hours(user, day), Decimal("0"), None
            hours = base_hours + overtime_hours
            total_hours += hours
            entry = entries_by_key.get((user.id, day))
            if entry is not None and entry.status == PRODUCTION_STATUS_PENDING:
                pending_count += 1
            if timesheet_entry is not None and timesheet_entry.status == PRODUCTION_STATUS_PENDING:
                pending_timesheet_count += 1
            days.append(SubordinateDay(
                date=day,
                hours=hours,
                base_hours=base_hours,
                overtime_hours=overtime_hours,
                text_value=text_value,
                status=entry.status if entry else None,
                entry_id=entry.id if entry else None,
                timesheet_status=timesheet_entry.status if timesheet_entry else None,
                timesheet_entry_id=timesheet_entry.id if timesheet_entry else None,
            ))
        summaries.append(SubordinateSummary(
            user=EmployeeBrief.model_validate(user),
            days=days,
            total_hours=total_hours,
            pending_count=pending_count,
            pending_timesheet_count=pending_timesheet_count,
        ))
    return summaries


@router.get("/subordinates/summary/export")
async def export_subordinates_summary_excel(
    year: int,
    month: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Выгружает в Excel сводный табель по всем подчинённым текущего
    пользователя за месяц - по строкам сотрудники (с разбивкой на
    "основное"/"сверхурочное" время), в конце строки - общее количество
    сданных приборов за месяц. Начальник цеха и выше попадает в выгрузку и
    своей собственной строкой.
    """
    from openpyxl import Workbook

    subordinates = await get_subordinates(db, current_user, include_self=True)
    if not subordinates:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No subordinates found")

    sub_ids = [u.id for u in subordinates]
    start, end = _month_bounds(year, month)

    hours_result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id.in_(sub_ids),
            TimesheetEntryModel.date >= start,
            TimesheetEntryModel.date <= end,
        )
    )
    hours_by_user: dict[int, dict[date, tuple[Decimal, Decimal, str | None]]] = {uid: {} for uid in sub_ids}
    for e in hours_result.scalars().all():
        hours_by_user[e.user_id][e.date] = (e.hours, e.overtime_hours, e.text_value)

    entries_result = await db.execute(
        select(ProductionEntryModel)
        .options(selectinload(ProductionEntryModel.items))
        .where(
            ProductionEntryModel.user_id.in_(sub_ids),
            ProductionEntryModel.date >= start,
            ProductionEntryModel.date <= end,
        )
    )
    products_by_user: dict[int, Decimal] = {uid: Decimal("0") for uid in sub_ids}
    for e in entries_result.scalars().all():
        products_by_user[e.user_id] += sum((item.quantity for item in e.items), Decimal("0"))

    rows = [
        {
            "number": idx + 1,
            "tab_number": user.id,
            "grade": get_table_role_label(user.role),
            "fio": _user_name(user),
            "hours_by_date": hours_by_user.get(user.id, {}),
            "products_total": products_by_user.get(user.id, Decimal("0")),
            "user": user,
        }
        for idx, user in enumerate(subordinates)
    ]

    workshop_label = await _get_workshop_label(db, current_user.workshop_id)

    wb = Workbook()
    ws_timesheet = wb.active
    ws_timesheet.title = "Табель"
    table_start_row = _write_export_summary_header(ws_timesheet, workshop_label, year, month, rows)
    _fill_timesheet_grid_sheet(ws_timesheet, year, month, rows, start_row=table_start_row)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"timesheet_subordinates_{year}_{month:02d}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/production/entries/pending", response_model=list[ProductionEntrySchema])
async def list_pending_entries(
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Записи, ожидающие согласования, по всем подчинённым текущего пользователя
    согласно иерархии должностей.
    """
    result = await db.execute(
        select(ProductionEntryModel)
        .options(
            selectinload(ProductionEntryModel.items).selectinload(ProductionEntryItemModel.product),
            selectinload(ProductionEntryModel.user),
        )
        .where(ProductionEntryModel.status == PRODUCTION_STATUS_PENDING)
        .order_by(ProductionEntryModel.submitted_at.asc())
    )
    entries = result.scalars().all()
    visible = [e for e in entries if can_review_production(current_user, e.user)]

    schemas = []
    for entry in visible:
        base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)
        schemas.append(_entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value))
    return schemas


@router.get("/production/{user_id}", response_model=list[DayTimesheetInfo])
async def get_user_production_month(
    user_id: int,
    year: int,
    month: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Данные вкладки "Количество сданной продукции" по дням месяца: часы из
    табеля (автоподгрузка) + запись о сданной продукции за каждый день, если
    она есть.
    """
    target = await _get_user_or_404(db, user_id)
    if not can_view_timesheet(current_user, target):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to this timesheet")

    start, end = _month_bounds(year, month)

    hours_result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id == user_id,
            TimesheetEntryModel.date >= start,
            TimesheetEntryModel.date <= end,
        )
    )
    hours_by_date = {e.date: e for e in hours_result.scalars().all()}

    entries_result = await db.execute(
        select(ProductionEntryModel)
        .options(
            selectinload(ProductionEntryModel.items).selectinload(ProductionEntryItemModel.product),
            selectinload(ProductionEntryModel.user),
            selectinload(ProductionEntryModel.reviewer),
        )
        .where(
            ProductionEntryModel.user_id == user_id,
            ProductionEntryModel.date >= start,
            ProductionEntryModel.date <= end,
        )
    )
    entries_by_date = {e.date: e for e in entries_result.scalars().all()}

    days = []
    for day_num in range(1, end.day + 1):
        day = date(year, month, day_num)
        timesheet_entry = hours_by_date.get(day)
        if timesheet_entry is not None:
            base_hours = timesheet_entry.hours
            overtime_hours = timesheet_entry.overtime_hours
            text_value = timesheet_entry.text_value
        else:
            base_hours, overtime_hours, text_value = _default_base_hours(target, day), Decimal("0"), None
        entry = entries_by_date.get(day)
        entry_schema = _entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value) if entry else None
        days.append(DayTimesheetInfo(
            date=day, hours=base_hours + overtime_hours,
            base_hours=base_hours, overtime_hours=overtime_hours, text_value=text_value, entry=entry_schema,
            timesheet_status=timesheet_entry.status if timesheet_entry else None,
            timesheet_entry_id=timesheet_entry.id if timesheet_entry else None,
        ))
    return days


@router.get("/production/entry/{entry_id}", response_model=ProductionEntryDetailSchema)
async def get_production_entry(
    entry_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    entry = await _load_entry(db, entry_id)
    if not (can_view_timesheet(current_user, entry.user)):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")

    base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)
    schema = ProductionEntryDetailSchema.model_validate(entry)
    schema.base_hours = base_hours
    schema.overtime_hours = overtime_hours
    schema.hours = base_hours + overtime_hours
    schema.text_value = text_value
    schema.can_edit = can_edit_production_entry(current_user, entry.user, entry.status)
    return schema


@router.post("/production", response_model=ProductionEntrySchema, status_code=status.HTTP_201_CREATED)
async def upsert_production_entry(
    data: ProductionEntryUpsert,
    submit: bool = False,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Создаёт или обновляет запись о сданной продукции сотрудника за день
    (всегда для самого текущего пользователя - правило 4). Если submit=True,
    запись сразу отправляется на согласование (правило 11 - нельзя отправить
    пустую запись) и уведомляется ближайший вышестоящий пользователь.
    """
    result = await db.execute(
        select(ProductionEntryModel)
        .options(selectinload(ProductionEntryModel.items))
        .where(
            ProductionEntryModel.user_id == current_user.id,
            ProductionEntryModel.date == data.date,
        )
    )
    entry = result.scalar_one_or_none()

    is_new = entry is None
    if entry is None:
        entry = ProductionEntryModel(
            user_id=current_user.id,
            date=data.date,
            status=PRODUCTION_STATUS_DRAFT,
        )
        db.add(entry)
        await db.flush()
    else:
        if not can_edit_production_entry(current_user, entry.user or current_user, entry.status):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Cannot edit an approved or rejected production entry",
            )
        for item in list(entry.items):
            await db.delete(item)
        await db.flush()

    old_values = None if is_new else {"items": [
        {"product_id": i.product_id, "quantity": float(i.quantity)} for i in entry.items
    ]}

    for item_data in data.items:
        db.add(ProductionEntryItemModel(
            entry_id=entry.id,
            product_id=item_data.product_id,
            quantity=item_data.quantity,
            comment=item_data.comment,
        ))

    was_rejected = entry.status == PRODUCTION_STATUS_REJECTED

    if submit:
        if not data.items:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot submit an empty entry")
        entry.status = PRODUCTION_STATUS_PENDING
        entry.submitted_at = datetime.now(timezone.utc)
        entry.reviewer_comment = None
        entry.reviewed_at = None

    await db.flush()

    log = ProductionAuditLogModel(
        entry_id=entry.id,
        actor_id=current_user.id,
        action="created" if is_new else ("resubmitted" if was_rejected and submit else "updated"),
        old_values=old_values,
        new_values={"items": [
            {"product_id": i.product_id, "quantity": float(i.quantity)} for i in data.items
        ], "submitted": submit},
    )
    db.add(log)
    await db.commit()

    entry = await _load_entry(db, entry.id)
    base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)

    if submit:
        await _notify_supervisors(db, entry, base_hours + overtime_hours, resubmit=was_rejected)

    return _entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value)


@router.post("/production/entry/{entry_id}/submit", response_model=ProductionEntrySchema)
async def submit_production_entry(
    entry_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """Отправляет уже существующую (черновую/отклонённую) запись на согласование."""
    entry = await _load_entry(db, entry_id)
    if not can_edit_production_entry(current_user, entry.user, entry.status):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Cannot submit this entry")
    if not entry.items:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot submit an empty entry")

    was_rejected = entry.status == PRODUCTION_STATUS_REJECTED
    entry.status = PRODUCTION_STATUS_PENDING
    entry.submitted_at = datetime.now(timezone.utc)
    entry.reviewer_comment = None
    entry.reviewed_at = None

    db.add(ProductionAuditLogModel(
        entry_id=entry.id,
        actor_id=current_user.id,
        action="resubmitted" if was_rejected else "submitted",
    ))
    await db.commit()

    entry = await _load_entry(db, entry.id)
    base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)
    await _notify_supervisors(db, entry, base_hours + overtime_hours, resubmit=was_rejected)
    return _entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value)


@router.post("/production/entry/{entry_id}/approve", response_model=ProductionEntrySchema)
async def approve_production_entry(
    entry_id: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """Подтверждает запись (правило 11 - только вышестоящий пользователь, не сам сотрудник)."""
    entry = await _load_entry(db, entry_id)
    if not can_review_production(current_user, entry.user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if entry.status != PRODUCTION_STATUS_PENDING:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Entry is not pending approval")

    entry.status = PRODUCTION_STATUS_APPROVED
    entry.reviewer_id = current_user.id
    entry.reviewed_at = datetime.now(timezone.utc)
    entry.reviewer_comment = None

    db.add(ProductionAuditLogModel(entry_id=entry.id, actor_id=current_user.id, action="approved"))
    await db.commit()

    entry = await _load_entry(db, entry.id)
    base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)

    message = {
        "type": "production_entry_approved",
        "entry_id": entry.id,
        "date": entry.date.isoformat(),
    }
    asyncio.create_task(notification_manager.send_personal_message(message, entry.user_id))
    if entry.user.fcm_token:
        asyncio.create_task(send_fcm_notification(
            token=entry.user.fcm_token,
            title="Запись о продукции подтверждена",
            body=f"Запись за {entry.date.isoformat()} подтверждена",
            data=message,
            sender_id=current_user.id,
        ))

    return _entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value)


@router.post("/production/entry/{entry_id}/reject", response_model=ProductionEntrySchema)
async def reject_production_entry(
    entry_id: int,
    data: ProductionEntryReject,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """Отклоняет запись - комментарий обязателен (правило 11)."""
    entry = await _load_entry(db, entry_id)
    if not can_review_production(current_user, entry.user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if entry.status != PRODUCTION_STATUS_PENDING:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Entry is not pending approval")

    entry.status = PRODUCTION_STATUS_REJECTED
    entry.reviewer_id = current_user.id
    entry.reviewed_at = datetime.now(timezone.utc)
    entry.reviewer_comment = data.comment

    db.add(ProductionAuditLogModel(
        entry_id=entry.id, actor_id=current_user.id, action="rejected", comment=data.comment
    ))
    await db.commit()

    entry = await _load_entry(db, entry.id)
    base_hours, overtime_hours, text_value = await _get_hours_for_date(db, entry.user, entry.date)

    message = {
        "type": "production_entry_rejected",
        "entry_id": entry.id,
        "date": entry.date.isoformat(),
        "comment": data.comment,
    }
    asyncio.create_task(notification_manager.send_personal_message(message, entry.user_id))
    if entry.user.fcm_token:
        asyncio.create_task(send_fcm_notification(
            token=entry.user.fcm_token,
            title="Запись о продукции отклонена",
            body=data.comment,
            data=message,
            sender_id=current_user.id,
        ))

    return _entry_to_schema(entry, current_user, base_hours, overtime_hours, text_value)


# ------------------------------------------------------------------
# Экспорт в Excel
# ------------------------------------------------------------------

@router.get("/{user_id}/export")
async def export_timesheet_excel(
    user_id: int,
    year: int,
    month: int,
    current_user: UserModel = Depends(get_current_production_user),
    db: AsyncSession = Depends(get_async_db),
):
    """
    Выгружает табель пользователя за месяц в Excel - один лист "Табель" в
    том же формате сводной сетки, что и выгрузка таблицы подчинённых
    (/timesheet/subordinates/summary/export), чтобы выгрузка только своего
    табеля выглядела так же, как строка в таблице подчинённых.
    """
    from openpyxl import Workbook

    target = await _get_user_or_404(db, user_id)
    if not can_view_timesheet(current_user, target):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to this timesheet")

    start, end = _month_bounds(year, month)

    hours_result = await db.execute(
        select(TimesheetEntryModel).where(
            TimesheetEntryModel.user_id == user_id,
            TimesheetEntryModel.date >= start,
            TimesheetEntryModel.date <= end,
        )
    )
    hours_by_date = {e.date: (e.hours, e.overtime_hours, e.text_value) for e in hours_result.scalars().all()}

    entries_result = await db.execute(
        select(ProductionEntryModel)
        .options(selectinload(ProductionEntryModel.items).selectinload(ProductionEntryItemModel.product))
        .where(
            ProductionEntryModel.user_id == user_id,
            ProductionEntryModel.date >= start,
            ProductionEntryModel.date <= end,
        )
    )
    entries_by_date = {e.date: e for e in entries_result.scalars().all()}

    products_total = sum(
        (item.quantity for entry in entries_by_date.values() for item in entry.items),
        Decimal("0"),
    )

    rows = [{
        "number": 1,
        "tab_number": target.id,
        "grade": get_table_role_label(target.role),
        "fio": _user_name(target),
        "hours_by_date": hours_by_date,
        "products_total": products_total,
        "user": target,
    }]
    workshop_label = await _get_workshop_label(db, target.workshop_id)

    wb = Workbook()
    ws_timesheet = wb.active
    ws_timesheet.title = "Табель"
    table_start_row = _write_export_summary_header(ws_timesheet, workshop_label, year, month, rows)
    _fill_timesheet_grid_sheet(ws_timesheet, year, month, rows, start_row=table_start_row)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    employee_name = _user_name(target)
    filename = f"timesheet_{employee_name or target.id}_{year}_{month:02d}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
